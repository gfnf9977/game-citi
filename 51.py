import openpyxl
import json
from openpyxl.styles import PatternFill
from tqdm import tqdm

FILL_GREEN = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
FILL_YELLOW = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
FILL_ORANGE = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
FILL_BLUE = PatternFill(start_color='9DE1FE', end_color='9DE1FE', fill_type='solid')


def clean_text(val):
    if val is None: return ""
    return str(val).strip().replace('\xa0', ' ').replace("'", "’").replace("`", "’")


def process_excel(excel_path, json_path):
    print("Завантаження еталонної бази JSON...")
    with open(json_path, 'r', encoding='utf-8') as f:
        ref_data = json.load(f)

    print(f"Відкриваємо Excel файл {excel_path}...")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}

    col_katotth = headers.get('katotth')
    col_name = headers.get('name:uk')
    col_old_name = headers.get('old_name:uk')
    col_region = headers.get('region')
    col_rayon = headers.get('rayon')
    col_hromada = headers.get('hromada')

    # Шукаємо колонку type:uk, якщо немає - створюємо нову
    col_type_uk = headers.get('type:uk')
    if not col_type_uk:
        col_type_uk = ws.max_column + 1
        ws.cell(row=1, column=col_type_uk).value = 'type:uk'

    excel_katotths = set()

    # 1. Оновлення існуючих рядків
    for row in tqdm(range(2, ws.max_row + 1), desc="Перевірка існуючих", unit=" рядків"):
        katotth_val = ws.cell(row=row, column=col_katotth).value
        if not katotth_val: continue

        katotth = clean_text(katotth_val)
        excel_katotths.add(katotth)

        if katotth in ref_data:
            ref = ref_data[katotth]
            changed_cells = []

            # Перевірка назви
            current_name_raw = ws.cell(row=row, column=col_name).value
            current_name = clean_text(current_name_raw)
            ref_name = clean_text(ref['name:uk'])

            if current_name and current_name != ref_name:
                current_old = ws.cell(row=row, column=col_old_name).value
                new_old = f"{current_old}; {current_name_raw}" if current_old else current_name_raw
                ws.cell(row=row, column=col_old_name).value = new_old
                ws.cell(row=row, column=col_name).value = ref['name:uk']
                changed_cells.extend([col_old_name, col_name])

            # Записуємо український тип поселення (місто/село/селище)
            current_type_uk = clean_text(ws.cell(row=row, column=col_type_uk).value)
            if current_type_uk != ref['type:uk']:
                ws.cell(row=row, column=col_type_uk).value = ref['type:uk']
                changed_cells.append(col_type_uk)

            # Перевірка регіону, району, громади (тепер пряме порівняння, бо назви повні)
            for col_idx, key in [(col_region, 'region'), (col_rayon, 'rayon'), (col_hromada, 'hromada')]:
                current_val = clean_text(ws.cell(row=row, column=col_idx).value)
                ref_val = clean_text(ref[key])

                # Записуємо, якщо значення відрізняється АБО якщо клітинка була пуста
                if current_val != ref_val:
                    ws.cell(row=row, column=col_idx).value = ref[key]
                    changed_cells.append(col_idx)

            # Фарбування старих рядків
            row_changed = len(changed_cells) > 0
            base_fill = FILL_YELLOW if row_changed else FILL_GREEN
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = base_fill
            for col_idx in changed_cells:
                ws.cell(row=row, column=col_idx).fill = FILL_ORANGE

    # 2. Додавання нових рядків
    missing_settlements = [(k_id, ref) for k_id, ref in ref_data.items() if k_id not in excel_katotths]

    if missing_settlements:
        start_new_row = ws.max_row + 1
        for i, (k_id, ref) in enumerate(tqdm(missing_settlements, desc="Додавання нових", unit=" рядків")):
            new_row_idx = start_new_row + i

            ws.cell(row=new_row_idx, column=col_katotth).value = k_id
            ws.cell(row=new_row_idx, column=col_type_uk).value = ref['type:uk']
            ws.cell(row=new_row_idx, column=col_name).value = ref['name:uk']
            ws.cell(row=new_row_idx, column=col_region).value = ref['region']
            ws.cell(row=new_row_idx, column=col_rayon).value = ref['rayon']
            ws.cell(row=new_row_idx, column=col_hromada).value = ref['hromada']

            for col in range(1, ws.max_column + 1):
                ws.cell(row=new_row_idx, column=col).fill = FILL_BLUE

    output_filename = f"updated_FINAL_{excel_path}"
    wb.save(output_filename)
    print(f"\nГотово! Збережено у файл: {output_filename}")


# Запуск на початковому файлі
process_excel('ua-name-places.xlsx', 'ukraine_reference_full.json')